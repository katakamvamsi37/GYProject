import csv
import io
from datetime import date
from openpyxl import load_workbook
from core.models import Expense, Member
from .base import APITestCase


class ExportTests(APITestCase):
    def test_authenticated_year_filter_and_stable_empty_headers(self):
        year = date.today().year
        Expense.objects.create(description="Current", amount=10, spent_on=date(year, 9, 1))
        Expense.objects.create(description="Previous", amount=20, spent_on=date(year - 1, 9, 1))
        response = self.client.get(f"/api/export/expenses/?year={year}")
        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        self.assertIn("Current", text)
        self.assertNotIn("Previous", text)
        empty = self.client.get(f"/api/export/expenses/?year={year - 2}").content.decode(
            "utf-8-sig"
        )
        self.assertEqual(empty.splitlines()[0], text.splitlines()[0])
        self.use(None)
        self.assertEqual(self.client.get("/api/export/expenses/").status_code, 401)

    def test_formula_injection_escaped_in_csv_and_xlsx(self):
        Member.objects.create(name="=1+1", email="formula@example.test")
        text = self.client.get("/api/export/members/").content.decode("utf-8-sig")
        self.assertIn("'=1+1", text)
        response = self.client.get("/api/export/members/?file_format=xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.active["A2"].value, "'=1+1")
        self.assertEqual(workbook.active["A2"].data_type, "s")

    def test_export_matches_status_search_and_category(self):
        year = date.today().year
        for description, status, category in [
            ("Matched", "approved", "Food"),
            ("Pending", "pending", "Food"),
            ("Other", "approved", "Other"),
        ]:
            Expense.objects.create(
                description=description,
                status=status,
                category=category,
                amount=1,
                spent_on=date(year, 9, 1),
            )
        response = self.client.get("/api/export/expenses/?status=approved&category=Food")
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual([row["description"] for row in rows], ["Matched"])

    def test_xlsx_monetary_cells_are_numeric(self):
        year = date.today().year
        Expense.objects.create(
            description="Numeric amount", amount="150.25", spent_on=date(year, 9, 1)
        )
        response = self.client.get("/api/export/expenses/?file_format=xlsx")
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.active["D2"].value, 150.25)
        self.assertEqual(workbook.active["D2"].data_type, "n")

    def test_budget_export_uses_approved_expense_total(self):
        from core.models import Plan

        year = date.today().year
        plan = Plan.objects.create(
            title="Export budget",
            budget=500,
            spent=999,
            category="Food",
            target_date=date(year, 9, 1),
        )
        Expense.objects.create(
            description="Approved",
            amount=50,
            status="approved",
            spent_on=date(year, 9, 1),
            plan=plan,
        )
        response = self.client.get("/api/export/plans/")
        self.assertEqual(response.status_code, 200)
        rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(float(rows[0]["actual_spent"]), 50)
